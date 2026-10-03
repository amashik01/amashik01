"""data/tiktok_topads.jsonl + data/tt_run.log -> pinkwear-womens-shoes-tiktok.xlsx"""
import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

FOOTWEAR = "label_22109000000"   # industry label observed on every BD footwear ad
rows = [json.loads(l) for l in open("data/tiktok_topads.jsonl", encoding="utf-8")]
foot = sorted([r for r in rows if r["industry_key"] == FOOTWEAR], key=lambda r: -r["like"])
other = [r for r in rows if r["industry_key"] != FOOTWEAR]
link = lambda r: f"https://ads.tiktok.com/business/creativecenter/topads/{r['id']}/pc/en"
PRODUCT = [("Sandal", r"sandal|স্যান্ডেল|স্যান্ডাল"), ("Leather shoe", r"leather|চামড়া"), ("Heel", r"heel|হিল"),
           ("Wholesale/China factory", r"wholesale|factory|supplier")]
def product(t):
    return ", ".join(n for n, rx in PRODUCT if re.search(rx, t, re.I)) or "জুতা (সাধারণ)"
def price(t):
    m = re.findall(r"[৳]\s*([0-9০-৯,]+)", t)
    return ", ".join(m)
def gender(t):
    return "মহিলাদের জুতা স্পষ্টভাবে লেখা নেই (সম্ভবত ইউনিসেক্স/ছেলেদের)"

H = PatternFill("solid", fgColor="1F3A5F")
wb = Workbook(); wb.remove(wb.active)
def sheet(name, head, data, widths):
    ws = wb.create_sheet(name); ws.append(head)
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = H; c.alignment = Alignment(wrap_text=True)
    for d in data: ws.append(d)
    for i, w in enumerate(widths): ws.column_dimensions[chr(65+i)].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for r in ws.iter_rows(min_row=2):
        for c in r: c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws

summary = [
 ["Pinkwear — মহিলাদের জুতা: TikTok রিসার্চ (বাংলাদেশ) | তারিখ: ৩ অক্টোবর ২০২৬"],
 [""],
 ["১. কী করা হয়েছে"],
 [f"TikTok Creative Center → Top Ads (দেশ: Bangladesh, গত ১৮০ দিন) পাবলিক পেজে ৪১টি কিওয়ার্ড (ইংরেজি, বাংলা, প্রতিযোগীদের নাম) দিয়ে সার্চ করা হয়েছে। মোট {len(rows)}টি অ্যাড পাওয়া গেছে, কিন্তু অনেক কিওয়ার্ড (fashion, comfort, city shop, cod) অপ্রাসঙ্গিক অ্যাডও এনেছে। জুতার ইন্ডাস্ট্রি লেবেলে আছে মাত্র {len(foot)}টি অ্যাড — এগুলো 'Videos' শিটে। বাকি {len(other)}টি 'Excluded' শিটে।"],
 [""],
 ["২. মূল ফলাফল"],
 ["Top Ads-এ বাংলাদেশের জুতার অ্যাড খুবই কম (১২টি), এবং এর মধ্যে মহিলাদের জুতা স্পষ্টভাবে বিক্রি করছে এমন অ্যাড প্রায় নেই। পাওয়া অ্যাডগুলো: (ক) Flash Leather — অরিজিনাল চামড়ার জুতা, 'রেক্সিনের দামে', 'অর্ডার কনফার্ম করতে ১ টাকাও অগ্রিম লাগবে না' (একাধিক ভার্সন, ছেলেদের লেদার জুতা বলে মনে হয়); (খ) 'Trendy Arizona Sandal' — ৳৮৯০–৯৯৯ 'Hot Offer' টেমপ্লেট, ইউনিসেক্স; (গ) প্রিমিয়াম চামড়ার জুতা — ২.২ লক্ষ লাইক, তবে ছেলেদের জুতার সম্ভাবনা বেশি; (ঘ) চায়না সেকেন্ডহ্যান্ড/হোলসেল শু।"],
 ["Facebook তুলনা: Facebook-এ হিল ৪৪%, স্যান্ডেল ২৩%, ফ্ল্যাট ২০% — অর্থাৎ মহিলাদের জুতার প্রতিযোগীরা TikTok Top Ads-এ প্রায় অনুপস্থিত। তারা TikTok-এ পেইড অ্যাড চালায় না বা Top Ads তালিকায় আসেনি। এটা Pinkwear-এর জন্য সুযোগ: TikTok-এ মহিলাদের জুতার পেইড অ্যাড প্রতিযোগিতা খুব কম।"],
 [""],
 ["৩. যা কাজ করতে দেখা গেছে (পাওয়া অ্যাড থেকে)"],
 ["• 'এক টাকাও অগ্রিম নয়' (COD) — Flash Leather-এর একাধিক অ্যাডে, CTR ০.২৪–০.৬১%।"],
 ["• দাম সরাসরি অ্যাডে (৳৮৯০/৮৯৯/৯৯৯ 'Hot Offer') — ইমোজি ✅ বুলেট, 'অফিস/ক্যাজুয়াল' ব্যবহার, বাংলা ও ইংরেজি দুই ভার্সন।"],
 ["• একই প্রোডাক্টের একাধিক ভার্সন একসাথে চালানো (Arizona Sandal ৪টি ভার্সন)।"],
 ["• Facebook-এ আমরা যা দেখেছি (COD, রেডি স্টক, ওয়ারেন্টি) তা TikTok কপিতেও কাজ করছে।"],
 [""],
 ["৪. Pinkwear-এর জন্য সুপারিশ"],
 ["• TikTok-এ এখন মহিলাদের জুতার ফাঁকা জায়গা বেশি। ১৫–৩০ সেকেন্ডের ট্রাই-অন/পায়ে হাঁটা ভিডিও + ৳১২৯৯–১৯৯৯ দাম + 'ক্যাশ অন ডেলিভারি, রেডি স্টক' দিয়ে শুরু করুন।"],
 ["• প্রথম ৩ সেকেন্ডে সমস্যা-হুক: 'অফিসে সারাদিন হিল পরে পা ব্যথা?', ফ্যাক্টরি প্রসেস ভিডিও, আনবক্সিং।"],
 ["• ৩টি ক্রিয়েটিভ ভার্সন একসাথে টেস্ট করুন (বাংলা/ইংরেজি কপি, ভিন্ন হুক)।"],
 [""],
 ["৫. সীমাবদ্ধতা (সৎভাবে)"],
 ["• Top Ads একটি কিউরেটেড তালিকা — সব অ্যাড নয়। বাংলাদেশের জন্য এটি খুবই ছোট।"],
 ["• TikTok-এর সাধারণ সার্চ ও প্রোফাইল পেজ লগইন ছাড়া ভিডিও দেখায় না ('Log in' ওয়াল)। আমি তা বাইপাস করিনি। তাই অর্গানিক ভিডিও, প্রতিযোগী অ্যাকাউন্টের ফলোয়ার ও ভিউ এই ফাইলে নেই।"],
 ["• ৫০ অ্যাকাউন্ট / ১৫০+ ভিডিওর লক্ষ্যমাত্রা পূরণ হয়নি। এটি পূরণ করতে আপনার লগইন করা Chrome-এ (Claude in Chrome) অর্গানিক সার্চ ও প্রতিযোগী প্রোফাইল ঘেঁটে দেখতে হবে।"],
 ["• Top Ads-এর দাম/বিক্রির তথ্য নেই; শুধু লাইক ও CTR (পরিসর) আছে। ভিডিও লিংকের মেয়াদ থাকে না, তাই স্থায়ী 'Ad link' দেওয়া হয়েছে।"],
]
ws = wb.create_sheet("সারসংক্ষেপ (বাংলা)")
for s in summary: ws.append(s)
ws.column_dimensions["A"].width = 150
for r in ws.iter_rows():
    for c in r: c.alignment = Alignment(wrap_text=True, vertical="top")
for r in (1, 3, 6, 10, 15, 20): ws.cell(r, 1).font = Font(bold=True, size=12)

sheet("Videos", ["Ad Link (permanent)", "Brand/Account", "Title / Hook / Copy", "Product", "Price (৳)", "Likes", "CTR %", "Duration (s)", "Objective", "Found via keyword", "Ad ID"],
      [[link(r), r["brand_name"] or "(নাম নেই)", r["ad_title"], product(r["ad_title"]), price(r["ad_title"]), r["like"], r.get("ctr"),
        round((r.get("video_info") or {}).get("duration") or 0, 1), r.get("objective_key", "").replace("campaign_objective_", ""), r.get("_kw"), r["id"]] for r in foot],
      [48, 20, 70, 24, 12, 10, 8, 10, 16, 16, 22])

log = open("data/tt_run.log", encoding="utf-8").read()
kws = re.findall(r"'(.+?)': total_count=(\w+) new=(\d+)", log)
sheet("Keywords", ["Keyword", "Results in Top Ads (BD, 180 days)", "New ads added"], [[k, t, n] for k, t, n in kws], [30, 30, 15])
sheet("Excluded", ["Ad Link", "Brand", "Title", "Industry label", "Keyword", "Reason"],
      [[link(r), r["brand_name"], r["ad_title"][:200], r["industry_key"], r.get("_kw"), "জুতার ইন্ডাস্ট্রি নয়"] for r in other], [48, 20, 70, 22, 16, 22])
wb.save("pinkwear-womens-shoes-tiktok.xlsx"); print("saved", len(foot), "footwear,", len(other), "excluded")
