"""data/fb_ads.jsonl -> facebook_womens_shoes_analysis.xlsx ; data/tiktok_ads.jsonl -> tiktok_womens_shoes_analysis.xlsx"""
import json, os, sys, time, re
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HDR = PatternFill("solid", fgColor="1F3A5F")


def load(p):
    return [json.loads(l) for l in open(p, encoding="utf-8")] if os.path.exists(p) else []


def sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, 22)
    return ws


def d(ts):
    return time.strftime("%Y-%m-%d", time.gmtime(ts)) if isinstance(ts, (int, float)) and ts else ""


def days(a):
    s, e = a.get("start_date"), a.get("end_date")
    if not isinstance(s, (int, float)): return ""
    e = e if isinstance(e, (int, float)) and e else time.time()
    return max(0, int((e - s) / 86400))


PRODUCT_TERMS = {"Heel": r"heel|হিল|hil", "Sandal": r"sandal|sandel|স্যান্ডেল", "Slipper": r"slipper|slipar|স্লিপার|chappal|চপ্পল",
                 "Flat/Ballerina": r"flat|ballerina|ফ্ল্যাট|বেলেরিনা", "Loafer": r"loafer|লোফার", "Bridal/Party": r"bridal|party|wedding|biye|বিয়ে|পার্টি",
                 "Kolhapuri/Khussa/Nagra": r"kolhapuri|khussa|nagra|কোলাপুরি|নাগরা", "Boots": r"boot|বুট", "Leather": r"leather|চামড়া"}
OFFER_TERMS = {"Cash on Delivery": r"cash on delivery|cod|ক্যাশ অন", "Free Delivery": r"free delivery|ফ্রি ডেলিভারি",
               "Discount/Offer": r"discount|offer|off|ছাড়|অফার", "Price shown": r"৳|tk|taka|টাকা|bdt", "Eid/Festival": r"eid|ঈদ|puja|পূজা"}


def build_fb(out="facebook_womens_shoes_analysis.xlsx"):
    ads = load("data/fb_ads.jsonl")
    by_ad = {}
    for a in ads:                       # dedupe across keywords
        x = by_ad.setdefault(a["ad_id"], {**a, "keywords": set()})
        x["keywords"].add(a["keyword"])
    ads = list(by_ad.values())
    pages = defaultdict(list)
    for a in ads: pages[a["page_id"] or a["page_name"]].append(a)

    wb = Workbook(); wb.remove(wb.active)
    prow = []
    for pid, L in sorted(pages.items(), key=lambda kv: -len(kv[1])):
        kws = set().union(*[a["keywords"] for a in L])
        longest = max(L, key=lambda a: days(a) if days(a) != "" else -1)
        prow.append([L[0]["page_name"], L[0]["page_url"], pid, len(L), sum(1 for a in L if a.get("is_active")),
                     len(kws), ", ".join(sorted(kws))[:500], days(longest), longest["ad_library_url"],
                     Counter(a["display_format"] for a in L).most_common(1)[0][0],
                     Counter(u for a in L for u in [a["landing_url"]] if u).most_common(1)[0][0] if any(a["landing_url"] for a in L) else ""])
    sheet(wb, "Page Analysis", ["Page Name", "Page Link", "Page ID", "Total Ads Found", "Active Ads", "Keywords Matched",
          "Keyword List", "Longest Ad Running (days)", "Longest Ad Link", "Main Format", "Top Landing URL"], prow, {"Keyword List": 50})

    arows = []
    for a in sorted(ads, key=lambda a: (a["page_name"], -(days(a) or 0) if days(a) != "" else 0)):
        arows.append([a["page_name"], a["page_url"], a["ad_id"], a["ad_library_url"], a["display_format"], d(a.get("start_date")),
                      d(a.get("end_date")), days(a), "Active" if a.get("is_active") else "Inactive", ", ".join(a.get("platforms") or []),
                      a["cta"], a["title"], a["ad_copy"][:1500], a["landing_url"],
                      "\n".join(a["image_urls"][:3]), "\n".join(a["video_urls"][:2]), ", ".join(sorted(a["keywords"]))[:300]])
    sheet(wb, "Ad Details", ["Page Name", "Page Link", "Ad ID", "Ad Link (permanent)", "Format", "Start", "End", "Days Running", "Status",
          "Platforms", "CTA", "Title", "Ad Copy", "Landing URL", "Image Link (CDN, may expire)", "Video Link (CDN, may expire)", "Keywords"],
          arows, {"Ad Copy": 60, "Ad Link (permanent)": 40})

    kw_ads, kw_pages = defaultdict(set), defaultdict(set)
    for a in load("data/fb_ads.jsonl"):
        kw_ads[a["keyword"]].add(a["ad_id"]); kw_pages[a["keyword"]].add(a["page_id"] or a["page_name"])
    done = [k.strip() for k in open("data/done_keywords.txt", encoding="utf-8")] if os.path.exists("data/done_keywords.txt") else []
    krows = sorted(([k, len(kw_ads.get(k, ())), len(kw_pages.get(k, ()))] for k in done), key=lambda r: -r[1])
    sheet(wb, "Keyword Report", ["Keyword", "Ads Found", "Pages Found"], krows, {"Keyword": 40})

    text = lambda a: f"{a['ad_copy']} {a['title']}".lower()
    mrows = [["Total unique ads", len(ads), ""], ["Total unique pages", len(pages), ""],
             ["Active ads", sum(1 for a in ads if a.get('is_active')), ""]]
    for f, n in Counter(a["display_format"] for a in ads).most_common(): mrows.append([f"Format: {f}", n, f"{n*100//max(1,len(ads))}%"])
    for name, rx in PRODUCT_TERMS.items():
        n = sum(1 for a in ads if re.search(rx, text(a), re.I)); mrows.append([f"Product mention: {name}", n, f"{n*100//max(1,len(ads))}%"])
    for name, rx in OFFER_TERMS.items():
        n = sum(1 for a in ads if re.search(rx, text(a), re.I)); mrows.append([f"Offer/Angle: {name}", n, f"{n*100//max(1,len(ads))}%"])
    for c, n in Counter(a["cta"] for a in ads if a["cta"]).most_common(8): mrows.append([f"CTA: {c}", n, ""])
    long = sorted([a for a in ads if days(a) != ""], key=days, reverse=True)[:30]
    sheet(wb, "Market Trend", ["Metric", "Count", "Share"], mrows, {"Metric": 45})
    sheet(wb, "Winning Ads (Top 30 by days)", ["Page", "Days Running", "Ad Link", "Format", "Ad Copy"],
          [[a["page_name"], days(a), a["ad_library_url"], a["display_format"], a["ad_copy"][:400]] for a in long], {"Ad Copy": 70, "Ad Link": 40})
    wb.save(out); print("saved", out, len(ads), "ads", len(pages), "pages")


def build_tt(out="tiktok_womens_shoes_analysis.xlsx"):
    rows = load("data/tiktok_ads.jsonl")
    wb = Workbook(); wb.remove(wb.active)
    sheet(wb, "TikTok Top Ads", ["Brand", "Title", "Ad Link (permanent)", "Likes", "Comments", "Shares", "CTR", "Keyword", "Cover (CDN)", "Video (CDN)"],
          [[r["brand"], r["title"], r["detail_url"], r["likes"], r["comments"], r["shares"], r["ctr"], r["keyword"], r["cover"], r["video"]] for r in rows],
          {"Title": 50, "Ad Link (permanent)": 45})
    wb.save(out); print("saved", out, len(rows))


if __name__ == "__main__":
    which = sys.argv[1] if len(sys.argv) > 1 else "all"
    if which in ("fb", "all"): build_fb()
    if which in ("tiktok", "all"): build_tt()
